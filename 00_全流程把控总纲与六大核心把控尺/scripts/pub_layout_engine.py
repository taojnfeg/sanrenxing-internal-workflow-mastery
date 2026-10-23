# -*- coding: utf-8 -*-
import re, docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def init_pub_doc():
    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11.0)
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)
    # Usable width = 8.5 - 1.3 = 7.2 inches
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal_style.font.size = Pt(10.5)
    return doc

def add_pub_p(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size_pt=11, font_cn="宋体", sb=4, sa=6, color=(20,20,20), first_line_indent_pt=0):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(sb)
    pf.space_after = Pt(sa)
    pf.line_spacing = 1.6
    if first_line_indent_pt:
        pf.first_line_indent = Pt(first_line_indent_pt)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size_pt)
    r.font.name = "Arial" if font_cn == "黑体" else "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font_cn)
    r.font.color.rgb = RGBColor(*color)
    return p

def set_cell_margins(cell, top=110, bottom=110, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, hex_color):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def add_pub_table(doc, headers, rows, col_widths_in=None, zebra=True):
    ncols = len(headers)
    if not col_widths_in:
        col_widths_in = [7.2 / ncols] * ncols
    t = doc.add_table(rows=1, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    # Header row
    hdr_row = t.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))
    trPr.append(OxmlElement('w:cantSplit'))
    for i, htxt in enumerate(headers):
        c = hdr_row.cells[i]
        c.width = Inches(col_widths_in[i])
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(c, top=130, bottom=130, left=150, right=150)
        set_cell_shading(c, "1F3864")
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        r = p.add_run(htxt)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.name = "Arial"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Data rows
    for r_idx, row_data in enumerate(rows):
        row = t.add_row()
        r_trPr = row._tr.get_or_add_trPr()
        r_trPr.append(OxmlElement('w:cantSplit'))
        bg_hex = "F7FAFC" if (zebra and r_idx % 2 == 1) else "FFFFFF"
        for i, val in enumerate(row_data):
            c = row.cells[i]
            c.width = Inches(col_widths_in[i])
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(c, top=105, bottom=105, left=145, right=145)
            if bg_hex != "FFFFFF":
                set_cell_shading(c, bg_hex)
            c.text = ""
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 and len(str(val)) <= 14) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.42
            r = p.add_run(str(val))
            r.bold = (i == 0)
            r.font.size = Pt(9.5)
            r.font.name = "Times New Roman"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            r.font.color.rgb = RGBColor(25, 25, 25)
    return t

def purify_process_jargon(txt: str) -> str:
    if not txt:
        return ""
    # Skip lines that are pure process titles or look registry headers
    if any(k in txt for k in [
        "已更至EP", "02_过程库/", "专属四宫格", "Look Registry", "全集服装造型注册表",
        "状态备案：承接前集", "状态备案：全剧24集完结归档", "渐进生长总则", "首集完成态基底",
        "正脸骨相特征：", "跨时空双态正面骨相", "正面骨相与防混淆死锁", "正面骨相与造型特征"
    ]):
        return ""
    if re.match(r"^人物\d+_", txt.strip()) or re.match(r"^场景\d+_", txt.strip()):
        return ""
    # Clean embedded process codes & timestamps
    txt = re.sub(r"•\s*出场场次与正脸铁证：", "• 场次与剧情轨迹：", txt)
    txt = re.sub(r"•\s*EP\d{2}\s*新增造型锁死：", "• 造型状态：", txt)
    txt = re.sub(r"•\s*EP\d{2}\s*行为与台词实证：", "• 剧情表现：", txt)
    txt = re.sub(r"•\s*EP\d{2}\s*登场场次：", "• 登场场次：", txt)
    txt = re.sub(r"\s*[CP]\d{2}(?:/[CP]\d{2})*\s*\[\d{2}:\d{2}(?:\.\d+)?\]", "", txt)
    txt = re.sub(r"\s*（?[CP]\d{2}(?:/[CP]\d{2})*\s*\[\d{2}:\d{2}(?:\.\d+)?\]）?", "", txt)
    txt = re.sub(r"，?正脸铁证\s*[CP]\d{2}[^，。）]*", "", txt)
    txt = re.sub(r"（?EP\d{2}\s*[\d:\.–\-]+\s*[CP]\d{2}[^）]*）?", "", txt)
    txt = re.sub(r"（?\d{2}:\d{2}\.\d+[–\-]\d{2}:\d{2}\.\d+[^）]*）?", "", txt)
    txt = re.sub(r"\b\d{2}:\d{2}\.\d+\b", "", txt)
    txt = re.sub(r"\b\d{2}(?:-\d)?\.mp4\b", "", txt)
    txt = re.sub(r"｜绑定\s*`?人脸库_\d+`?", "", txt)
    txt = re.sub(r"（绑定\s*`?《?人脸库_\d+》?`?）", "", txt)
    txt = re.sub(r"【防呆锁\s*[A-Z][^】]*】：?", "", txt)
    txt = re.sub(r"\s+[CP]\d{2}\b", "", txt)
    txt = txt.replace("`", "").replace("*", "").strip()
    return txt
